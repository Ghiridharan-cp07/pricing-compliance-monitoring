import os
import re
import numbers
from datetime import date, datetime
from decimal import Decimal

import pandas as pd
from docx import Document
from db_connection import get_db_connection


TEMPLATE_FILE = (
    "templates/"
    "Pricing_Compliance_Violation_Notification_Template.docx"
)
OUTPUT_FOLDER = "output"


def load_violations():
    connection = get_db_connection()

    query = """
        SELECT
            SKU,
            Violation_Date,
            Category,
            Sub_category,
            Region,
            Marketplace,
            Seller_Name,
            MAP,
            LPP,
            Advertised_Price,
            Season,
            Promotion_Price
        FROM VIOLATION
        ORDER BY Seller_Name, Violation_Date, SKU
    """

    try:
        df = pd.read_sql(query, connection)
    finally:
        connection.close()

    if df.empty:
        raise ValueError(
            "No violation records found in VIOLATION table."
        )

    df = df[df["Seller_Name"].notna()].copy()

    if df.empty:
        raise ValueError(
            "No violation records have a valid Seller_Name."
        )

    return df


def safe_filename(name):
    filename = re.sub(
        r'[\\/:*?"<>|]',
        "_",
        str(name)
    ).strip().rstrip(".")

    return filename or "Unknown_Seller"


def replace_placeholder(document, placeholder, value):
    value = str(value)

    def replace_in_paragraph(paragraph):
        # Normal case: placeholder is contained in one run.
        for run in paragraph.runs:
            if placeholder in run.text:
                run.text = run.text.replace(
                    placeholder,
                    value
                )
                return

        # Fallback: placeholder may span multiple Word runs.
        if placeholder in paragraph.text:
            updated_text = paragraph.text.replace(
                placeholder,
                value
            )

            if paragraph.runs:
                paragraph.runs[0].text = updated_text

                for run in paragraph.runs[1:]:
                    run.text = ""
            else:
                paragraph.add_run(updated_text)

    for paragraph in document.paragraphs:
        replace_in_paragraph(paragraph)

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    replace_in_paragraph(paragraph)


def build_violation_description(seller_df):
    has_promotion = (
        seller_df["Promotion_Price"].fillna(0).gt(0).any()
    )

    if has_promotion:
        return "The listing reflects promotional pricing but remains below the applicable LPP."
    return ""


def format_value(value):
    if pd.isna(value):
        return "N/A"

    if isinstance(value, float):
        return f"{value:,.2f}"

    return str(value)


def populate_violation_table(document, seller_df):
    if not document.tables:
        raise ValueError(
            "The template must contain a pricing table."
        )

    # The first table must be the Pricing Violations table.
    table = document.tables[0]

    # Preserve the header; remove the sample data rows.
    while len(table.rows) > 1:
        table._tbl.remove(table.rows[1]._tr)

    for _, row in seller_df.iterrows():
        cells = table.add_row().cells

        values = [
            ("SKU", row["SKU"]),
            ("Category", row["Category"]),
            ("Sub_category", row["Sub_category"]),
            ("Region", row["Region"]),
            ("Marketplace", row["Marketplace"]),
            ("Violation_Date", row["Violation_Date"]),
            ("MAP", row["MAP"]),
            ("LPP", row["LPP"]),
            ("Advertised_Price", row["Advertised_Price"]),
            ("Promotion_Price", row["Promotion_Price"])
        ]

        for cell, (column, value) in zip(cells, values):
            cell.text = format_value(value)


def create_seller_document(seller_name, seller_df):
    document = Document(TEMPLATE_FILE)

    seasons = seller_df["Season"].dropna()

    season = (
        str(seasons.iloc[0])
        if not seasons.empty
        else "Q4"
    )

    replacements = {
        "[SELLER_NAME]": seller_name,
        "[SEASON]": season,
        "[VIOLATION_DESCRIPTION]":
            build_violation_description(seller_df),
        "[Company Name]": "XYZ Solutions",
        "[Contact Information]":
            "xyz.compliance@example.com"
    }

    for placeholder, value in replacements.items():
        replace_placeholder(
            document,
            placeholder,
            value
        )

    populate_violation_table(
        document,
        seller_df
    )

    os.makedirs(OUTPUT_FOLDER, exist_ok=True)

    filename = (
        f"{safe_filename(seller_name)}_Violation.docx"
    )

    output_path = os.path.join(
        OUTPUT_FOLDER,
        filename
    )

    document.save(output_path)

    print(f"Created: {output_path}")


def main():
    violations = load_violations()

    total_records = len(violations)
    total_sellers = violations["Seller_Name"].nunique()

    print(f"Total violation records: {total_records}")
    print(f"Total seller files to create: {total_sellers}")
    print("-" * 50)

    processed_files = 0

    for seller_name, seller_df in violations.groupby(
        "Seller_Name",
        sort=True
    ):
        create_seller_document(
            seller_name,
            seller_df
        )

        processed_files += 1

        print(
            f"Files processed: "
            f"{processed_files}/{total_sellers}"
        )

    print("-" * 50)
    print(
        f"Completed. Total files created: "
        f"{processed_files}"
    )


if __name__ == "__main__":
    main()